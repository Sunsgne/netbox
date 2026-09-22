"""Load ZENLENET workbooks into the operations database.

Spreadsheets stay outside the git repo. Credential-looking cells are dropped
and never stored.
"""

from __future__ import annotations

import argparse
from datetime import date, datetime, timedelta

import openpyxl
from sqlalchemy import delete, func, select

from app.db import SessionLocal, init_db
from app.models import (
    Circuit,
    Customer,
    Invoice,
    InvoiceLine,
    IpRecord,
    ServiceOrder,
    SupplierReturn,
    Ticket,
    TicketEvent,
    VxlanLink,
)
from app.parsing import (
    classify_customer,
    clean_remark,
    dc_type_for,
    decide_status,
    extract_ips,
    extract_v6,
    first_line,
    is_private_v4,
    parse_bandwidth,
    parse_dates,
    parse_expiry,
    product_of,
    role_from_header,
)

GRID_SKIP = {
    "说明",
    "游戏列表",
    "邮件告警模板",
    "供应商接口",
    "专线端口",
    "SERVER业务记录",
    "剩余地址",
    "供应商专线记录",
    "Zenlayer所有IP",
    "电科310个美国IP业务记录",
}

STATUS_RANK = {
    "allocated": 5,
    "returning": 4,
    "testing": 3,
    "reserved": 2,
    "internal": 2,
    "free": 1,
}


def import_all(ip_path: str, company_path: str, ipv6_path: str) -> dict:
    init_db()
    with SessionLocal() as session:
        _reset(session)
        customers: dict[str, Customer] = {}
        pending: dict[tuple, dict] = {}
        _import_grids(ip_path, session, customers, pending)
        _import_ipv6(ipv6_path, session, customers, pending)
        _flush_ips(session, pending, customers)
        orders = _import_company(company_path, session, customers)
        _refresh_customers(session)
        tickets = _seed_tickets(session)
        session.commit()
        stats = {
            "customers": session.scalar(select(func.count()).select_from(Customer)) or 0,
            "addresses": session.scalar(select(func.count()).select_from(IpRecord)) or 0,
            "orders": orders,
            "circuits": session.scalar(select(func.count()).select_from(Circuit)) or 0,
            "tickets": tickets,
        }
    return stats


def _reset(session) -> None:
    for model in (
        InvoiceLine,
        Invoice,
        TicketEvent,
        Ticket,
        ServiceOrder,
        IpRecord,
        VxlanLink,
        Circuit,
        SupplierReturn,
        Customer,
    ):
        session.execute(delete(model))
    session.flush()


def _customer(session, cache: dict[str, Customer], name: str) -> Customer:
    found = cache.get(name)
    if found:
        return found
    found = session.scalar(select(Customer).where(Customer.name == name))
    if found:
        cache[name] = found
        return found
    found = Customer(name=name, kind="external", status="active")
    session.add(found)
    session.flush()
    cache[name] = found
    return found


def _consider(pending: dict, row: dict) -> None:
    key = (row["version"], row["address"], row["prefixlen"])
    current = pending.get(key)
    if current is None:
        pending[key] = row
        return
    new_rank = STATUS_RANK.get(row["status"], 0)
    old_rank = STATUS_RANK.get(current["status"], 0)
    if new_rank > old_rank or (row.get("customer_name") and not current.get("customer_name")):
        pending[key] = row


def _import_grids(path: str, session, customers, pending) -> None:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    try:
        for name in workbook.sheetnames:
            if name.strip() in GRID_SKIP:
                continue
            _import_grid(workbook[name], name.strip(), session, customers, pending)
    finally:
        workbook.close()


def _import_grid(ws, pop: str, session, customers, pending) -> None:
    header_row = None
    max_col = min(ws.max_column or 1, 80)
    max_row = ws.max_row or 1
    for row_idx in range(1, min(8, max_row) + 1):
        for col_idx in range(1, max_col + 1):
            value = ws.cell(row_idx, col_idx).value
            if isinstance(value, str) and value.strip().upper() == "ISP":
                header_row = row_idx
                break
        if header_row:
            break
    if not header_row:
        return
    isp_cols = []
    for col_idx in range(1, max_col + 1):
        value = ws.cell(header_row, col_idx).value
        if isinstance(value, str) and value.strip().upper() == "ISP":
            isp_cols.append(col_idx)
    groups = []
    for index, col_idx in enumerate(isp_cols):
        nxt = isp_cols[index + 1] if index + 1 < len(isp_cols) else col_idx + 4
        ip_header = ws.cell(header_row, col_idx + 1).value
        if not ip_header or "IP" not in str(ip_header).upper():
            continue
        groups.append(
            {
                "isp": col_idx,
                "ip": col_idx + 1,
                "cust": col_idx + 2 if col_idx + 2 < nxt else None,
                "remark": col_idx + 3 if col_idx + 3 < nxt else None,
                "role": role_from_header(str(ip_header)),
            }
        )
    today = date.today()
    for group in groups:
        supplier = ""
        for row_idx in range(header_row + 1, max_row + 1):
            supplier_cell = ws.cell(row_idx, group["isp"]).value
            supplier_text = first_line(supplier_cell)
            if supplier_text and len(supplier_text) <= 60 and "密码" not in supplier_text:
                supplier = supplier_text
            ip_cell = ws.cell(row_idx, group["ip"])
            cust_cell = ws.cell(row_idx, group["cust"]) if group["cust"] else None
            remark_cell = ws.cell(row_idx, group["remark"]) if group["remark"] else None
            targets = extract_ips(ip_cell.value) or extract_v6(ip_cell.value)
            if not targets:
                continue
            raw_customer = cust_cell.value if cust_cell is not None else None
            kind, cname = classify_customer(raw_customer)
            remark = clean_remark(
                remark_cell.value if remark_cell is not None else None,
                raw_customer if kind in {"skip", "infra"} else None,
            )
            rgb = _fill_rgb(cust_cell) or _fill_rgb(ip_cell)
            label = first_line(raw_customer)
            if kind == "hold":
                status, expiry = "reserved", parse_expiry(f"{label} {remark}")
            else:
                status, expiry = decide_status(kind, label, remark, rgb, today)
            usage = ""
            if kind == "infra":
                usage = label[:40]
            version = 6 if ":" in targets[0][0] else 4
            for address, plen in targets:
                net_attr = "内网" if version == 4 and is_private_v4(address) else "公网"
                _consider(
                    pending,
                    {
                        "version": version,
                        "address": address,
                        "prefixlen": plen,
                        "pop": pop,
                        "supplier": supplier[:80],
                        "customer_name": cname if kind in {"customer", "hold"} else None,
                        "role": group["role"],
                        "net_attr": net_attr,
                        "dc_type": dc_type_for(supplier),
                        "status": status,
                        "usage": usage,
                        "remark": remark,
                        "expires_on": expiry,
                        "source": pop,
                    },
                )


def _fill_rgb(cell) -> str:
    if cell is None or not cell.fill or cell.fill.patternType != "solid" or not cell.fill.fgColor:
        return ""
    color = cell.fill.fgColor
    if color.type == "rgb" and color.rgb:
        return str(color.rgb)[-6:].upper()
    return ""


def _import_ipv6(path: str, session, customers, pending) -> None:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for name in workbook.sheetnames:
            ws = workbook[name]
            rows = list(ws.iter_rows(values_only=True))
            if name.strip().upper() == "IPV6":
                for row in rows[1:]:
                    parent, child, country, origin_as, announce_as, _child2, customer = (list(row) + [None] * 7)[:7]
                    for addr, plen in extract_v6(child) or extract_v6(parent):
                        kind, cname = classify_customer(customer)
                        _consider(
                            pending,
                            {
                                "version": 6,
                                "address": addr,
                                "prefixlen": plen,
                                "pop": first_line(country) or "IPv6",
                                "supplier": "Zenlenet",
                                "customer_name": cname if kind in {"customer", "hold"} else None,
                                "role": "native",
                                "net_attr": "公网",
                                "dc_type": "主营机房",
                                "status": "reserved" if kind == "hold" else ("allocated" if kind == "customer" else "free"),
                                "usage": "",
                                "remark": clean_remark(f"parent {parent}" if parent else "", f"AS{origin_as}" if origin_as else "", f"代播 AS{announce_as}" if announce_as else ""),
                                "expires_on": None,
                                "source": "ipv6",
                            },
                        )
                continue
            header = None
            for index, row in enumerate(rows[:6]):
                if row and first_line(row[0]).upper() == "ISP":
                    header = index
                    break
            if header is None:
                continue
            supplier = ""
            for row in rows[header + 1 :]:
                values = list(row) + [None, None, None]
                if first_line(values[0]):
                    supplier = first_line(values[0])[:80]
                kind, cname = classify_customer(values[2])
                for addr, plen in extract_v6(values[1]):
                    _consider(
                        pending,
                        {
                            "version": 6,
                            "address": addr,
                            "prefixlen": plen,
                            "pop": name.strip(),
                            "supplier": supplier or "Zenlenet",
                            "customer_name": cname if kind in {"customer", "hold"} else None,
                            "role": "native",
                            "net_attr": "公网",
                            "dc_type": dc_type_for(supplier or "Zenlenet"),
                            "status": "reserved" if kind == "hold" else ("allocated" if kind == "customer" else "free"),
                            "usage": "",
                            "remark": clean_remark(values[0]),
                            "expires_on": None,
                            "source": "ipv6",
                        },
                    )
    finally:
        workbook.close()


def _flush_ips(session, pending: dict, customers: dict[str, Customer]) -> None:
    for row in pending.values():
        customer_id = None
        if row.get("customer_name"):
            customer_id = _customer(session, customers, row["customer_name"]).id
        session.add(
            IpRecord(
                version=row["version"],
                address=row["address"],
                prefixlen=row["prefixlen"],
                pop=row["pop"],
                supplier=row["supplier"],
                customer_id=customer_id,
                role=row["role"],
                net_attr=row["net_attr"],
                dc_type=row["dc_type"],
                status=row["status"],
                usage=row["usage"],
                remark=row["remark"],
                expires_on=row["expires_on"],
                source=row["source"],
            )
        )
    session.flush()


def _import_company(path: str, session, customers) -> int:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    count = 0
    try:
        sheets = {name.strip(): workbook[name] for name in workbook.sheetnames}
        count += _orders_from_rmipt(sheets.get("RM&IPT业务记录"), session, customers, "active", "rmipt")
        count += _orders_from_rmipt(sheets.get("RMIPT客户退租"), session, customers, "terminated", "rmipt-off")
        count += _orders_from_nanjie(sheets.get("尊领-南捷"), session, customers)
        count += _orders_from_pl(sheets.get("PL&SDWAN-业务记录"), session, customers)
        count += _orders_from_server(sheets.get("SERVER业务记录第二版"), session, customers, "active")
        count += _orders_from_server(sheets.get("客户测试"), session, customers, "testing")
        count += _orders_from_colo(sheets.get("托管业务记录"), session, customers)
        count += _orders_from_resale(sheets.get("纯转售"), session, customers)
        count += _orders_from_xc(sheets.get("XC线路"), session, customers)
        _vxlan(sheets.get("VXLAN记录"), session)
        _returns(sheets.get("退资源"), session)
        _circuits(sheets.get("供应商专线记录"), session)
    finally:
        workbook.close()
    return count


def _rows(ws):
    if ws is None:
        return []
    return [tuple(row) for row in ws.iter_rows(values_only=True)]


def _text(value) -> str:
    return clean_remark(value)


def _add_order(session, customers, **kwargs) -> bool:
    name = kwargs.pop("customer_name", None)
    if not name:
        return False
    payload = {key: kwargs.get(key) or "" for key in (
        "product", "status", "country", "pop_code", "bandwidth_text", "ip_text",
        "supplier_name", "circuit_no", "vlan", "device", "spec", "source", "note",
    )}
    payload["bw_mbps"] = kwargs.get("bw_mbps")
    payload["started_on"] = kwargs.get("started_on")
    payload["ended_on"] = kwargs.get("ended_on")
    payload["customer_id"] = _customer(session, customers, name).id
    if payload["status"] == "terminated" and payload["note"]:
        pass
    session.add(ServiceOrder(**payload))
    return True


def _orders_from_rmipt(ws, session, customers, status, source) -> int:
    rows = _rows(ws)
    if not rows:
        return 0
    header_at = next((i for i, row in enumerate(rows[:5]) if first_line(row[0]).lower() == "customer"), None)
    if header_at is None:
        return 0
    count = 0
    last = ""
    for row in rows[header_at + 1 :]:
        values = list(row) + [None] * 12
        name = first_line(values[0]) or last
        if first_line(values[0]):
            last = name
        kind, cname = classify_customer(name)
        if kind != "customer":
            continue
        ip_text = _text(values[6])
        bw_text = first_line(values[3])[:80]
        if not any((values[1], values[2], values[3], values[5], values[6])):
            continue
        if _add_order(
            session,
            customers,
            customer_name=cname,
            product=product_of(first_line(values[5]), "RMIPT"),
            status=status,
            country=first_line(values[1])[:80],
            pop_code=first_line(values[2])[:80],
            bandwidth_text=bw_text,
            bw_mbps=parse_bandwidth(values[3]),
            ip_text=ip_text,
            supplier_name=first_line(values[7])[:80],
            device=first_line(values[8])[:160],
            vlan=first_line(values[9])[:40],
            note=_text(values[4])[:300],
            source=source,
        ):
            count += 1
    return count


def _orders_from_nanjie(ws, session, customers) -> int:
    rows = _rows(ws)
    count = 0
    for row in rows[1:]:
        values = list(row) + [None] * 12
        if not any(values[:8]):
            continue
        if _add_order(
            session,
            customers,
            customer_name="南捷",
            product=product_of(first_line(values[4]), "RMIPT"),
            status="active",
            country=first_line(values[0])[:80],
            pop_code=first_line(values[1])[:80],
            bandwidth_text=first_line(values[2])[:80],
            bw_mbps=parse_bandwidth(values[2]),
            ip_text=_text(values[5]),
            supplier_name=first_line(values[6])[:80],
            device=first_line(values[7])[:160],
            vlan=first_line(values[8])[:40],
            note=_text(values[3])[:300],
            source="nanjie",
        ):
            count += 1
    return count


def _orders_from_pl(ws, session, customers) -> int:
    rows = _rows(ws)
    count = 0
    last = ""
    for row in rows[1:]:
        values = list(row) + [None] * 10
        if first_line(values[0]):
            last = first_line(values[0])
        kind, cname = classify_customer(last)
        if kind != "customer":
            continue
        if not any(values[1:6]):
            continue
        ended = parse_dates(values[9])
        status = "terminated" if ended and ended[-1] < date.today() else "active"
        ips = " ".join(bit for bit in (_text(values[4]), _text(values[5])) if bit)
        if _add_order(
            session,
            customers,
            customer_name=cname,
            product=product_of(first_line(values[1]), "PL"),
            status=status,
            country=first_line(values[2])[:80],
            pop_code=first_line(values[2])[:80],
            bandwidth_text=first_line(values[3])[:80],
            bw_mbps=parse_bandwidth(values[3]),
            ip_text=ips,
            supplier_name=first_line(values[6])[:80],
            circuit_no=first_line(values[7])[:80],
            started_on=parse_dates(values[8])[-1] if parse_dates(values[8]) else None,
            ended_on=ended[-1] if ended else None,
            source="pl",
        ):
            count += 1
    return count


def _orders_from_server(ws, session, customers, forced_status) -> int:
    rows = _rows(ws)
    if len(rows) < 2:
        return 0
    # 客户测试 has a banner row before the header.
    header_at = 0
    if first_line(rows[0][0]).lower() != "customer":
        header_at = 1 if len(rows) > 1 and first_line(rows[1][0]).lower() == "customer" else 0
    count = 0
    last = ""
    for row in rows[header_at + 1 :]:
        values = list(row) + [None] * 14
        if first_line(values[0]):
            last = first_line(values[0])
        kind, cname = classify_customer(last)
        if kind != "customer":
            continue
        if not any((values[2], values[6], values[8], values[3])):
            continue
        ended = parse_dates(values[12]) if len(values) > 12 else []
        if forced_status == "testing":
            status = "testing"
        elif ended and ended[-1] <= date.today():
            status = "terminated"
        elif values[12] and "退" in str(values[12]):
            status = "terminated"
        else:
            status = "active"
        if _add_order(
            session,
            customers,
            customer_name=cname,
            product=product_of(first_line(values[1]), "VM"),
            status=status,
            country=first_line(values[2])[:80],
            pop_code=first_line(values[2])[:80],
            bandwidth_text=first_line(values[3])[:80],
            bw_mbps=parse_bandwidth(values[3]),
            spec=first_line(values[6])[:80],
            ip_text=_text(values[8]),
            supplier_name=first_line(values[10])[:80],
            started_on=parse_dates(values[11])[-1] if parse_dates(values[11]) else None,
            ended_on=ended[-1] if ended else None,
            source="server" if forced_status != "testing" else "test",
        ):
            count += 1
    return count


def _orders_from_colo(ws, session, customers) -> int:
    rows = _rows(ws)
    count = 0
    last = ""
    country = ""
    dc = ""
    for row in rows[1:]:
        values = list(row) + [None] * 8
        if first_line(values[0]):
            last = first_line(values[0])
        if first_line(values[1]):
            country = first_line(values[1])
        if first_line(values[3]):
            dc = first_line(values[3])
        kind, cname = classify_customer(last)
        if kind != "customer" or not first_line(values[2]):
            continue
        if _add_order(
            session,
            customers,
            customer_name=cname,
            product="托管",
            status="active",
            country=country[:80],
            pop_code=(dc or country)[:80],
            spec=first_line(values[2])[:80],
            device=first_line(values[4])[:160],
            note=clean_remark(values[5], values[6]),
            source="colo",
        ):
            count += 1
    return count


def _orders_from_resale(ws, session, customers) -> int:
    rows = _rows(ws)
    count = 0
    for row in rows[1:]:
        values = list(row) + [None, None, None]
        kind, cname = classify_customer(values[1])
        if kind != "customer":
            continue
        note = _text(values[2])
        if _add_order(
            session,
            customers,
            customer_name=cname,
            product="转售",
            status="active",
            bandwidth_text=first_line(note)[:80],
            bw_mbps=parse_bandwidth(note),
            note=note,
            source="resale",
        ):
            count += 1
    return count


def _orders_from_xc(ws, session, customers) -> int:
    rows = _rows(ws)
    count = 0
    for row in rows[1:]:
        values = list(row) + [None] * 11
        kind, cname = classify_customer(values[0])
        if kind != "customer":
            continue
        blob = " ".join(str(item) for item in values if item)
        status = "terminated" if "退" in blob else "active"
        if _add_order(
            session,
            customers,
            customer_name=cname,
            product="PL",
            status=status,
            pop_code=first_line(values[1])[:80],
            country=first_line(values[6])[:80],
            device=first_line(values[2])[:160],
            circuit_no=first_line(values[7])[:80],
            note=clean_remark(values[10], values[8]),
            started_on=parse_dates(values[9])[-1] if parse_dates(values[9]) else None,
            source="xc",
        ):
            count += 1
    return count


def _vxlan(ws, session) -> None:
    for row in _rows(ws)[1:]:
        values = list(row) + [None] * 9
        if not any(values[:4]):
            continue
        session.add(
            VxlanLink(
                vni=first_line(values[0])[:20],
                a_end=first_line(values[1])[:160],
                z_end=first_line(values[2])[:160],
                customer_name=first_line(values[3])[:80],
                vlan=first_line(values[4])[:40],
                bandwidth=first_line(values[5])[:40],
                purpose=first_line(values[6])[:80],
                stopped=first_line(values[8])[:40],
            )
        )


def _returns(ws, session) -> None:
    for row in _rows(ws)[1:]:
        values = list(row) + [None] * 4
        if not first_line(values[0]) and not first_line(values[1]):
            continue
        session.add(
            SupplierReturn(
                supplier=first_line(values[0])[:80],
                resource=first_line(values[1])[:160],
                when_text=first_line(values[2])[:40],
                note=_text(values[3]),
            )
        )


def _circuits(ws, session) -> None:
    for row in _rows(ws)[1:]:
        values = list(row) + [None] * 9
        if not any((values[0], values[1], values[4])):
            continue
        try:
            bandwidth = float(values[7] or 0)
            peak = float(values[8] or 0)
        except (TypeError, ValueError):
            bandwidth, peak = 0, 0
        session.add(
            Circuit(
                circuit_no=first_line(values[0])[:80],
                a_city=first_line(values[1])[:40],
                a_vlan=first_line(values[2])[:40],
                z_city=first_line(values[4])[:40],
                z_vlan=first_line(values[5])[:40],
                bandwidth_mbps=bandwidth,
                peak_mbps=peak,
            )
        )


def _refresh_customers(session) -> None:
    active_customers = set(
        session.scalars(select(ServiceOrder.customer_id).where(ServiceOrder.status == "active"))
    )
    testing_customers = set(
        session.scalars(select(ServiceOrder.customer_id).where(ServiceOrder.status == "testing"))
    )
    allocated_customers = set(
        session.scalars(select(IpRecord.customer_id).where(IpRecord.status == "allocated", IpRecord.customer_id.is_not(None)))
    )
    for customer in session.scalars(select(Customer)):
        if customer.id in active_customers or customer.id in allocated_customers:
            customer.status = "active"
        elif customer.id in testing_customers:
            customer.status = "testing"
        else:
            customer.status = "churned"


def _seed_tickets(session) -> int:
    from app.content import TICKET_TYPES
    from app.models import NoticeTemplate
    from app.parsing import render_notice

    if session.scalar(select(func.count()).select_from(Ticket)):
        return 0
    ranked = session.execute(
        select(Customer, func.count(ServiceOrder.id))
        .join(ServiceOrder, ServiceOrder.customer_id == Customer.id)
        .where(ServiceOrder.status == "active")
        .group_by(Customer.id)
        .order_by(func.count(ServiceOrder.id).desc())
        .limit(2)
    ).all()
    if not ranked:
        return 0
    templates = {row.code: row for row in session.scalars(select(NoticeTemplate))}
    start = (datetime.utcnow() + timedelta(hours=8)).replace(minute=0, second=0, microsecond=0) + timedelta(days=2)
    end = start + timedelta(hours=2)
    created = 0
    samples = [
        ("cutover", "HKG", "163.53.245.0/26 及相关专线"),
        ("maintenance", "LAX", "LAX 接入节点与在网 IPT"),
    ]
    for (customer, _), (type_code, pop, impact) in zip(ranked, samples):
        code = dict((item[0], item[2]) for item in TICKET_TYPES)[type_code]
        template = templates.get(code)
        ctx = {
            "place": pop,
            "impact": impact,
            "start": start.strftime("%Y-%m-%d %H:%M"),
            "end": end.strftime("%Y-%m-%d %H:%M"),
            "utc_start": (start - timedelta(hours=8)).strftime("%Y-%m-%d %H:%M"),
            "utc_end": (end - timedelta(hours=8)).strftime("%Y-%m-%d %H:%M"),
            "reason": "网络出口设备及专线设备升级维护",
            "duration": "30分钟",
            "action": "线路割接" if type_code == "cutover" else "配置升级",
            "cause_by": "机房线路",
            "symptom": "短时中断",
        }
        subject = render_notice(template.subject if template else "", ctx)
        body = render_notice(template.body if template else "", ctx)
        title = "割接" if type_code == "cutover" else "维护"
        ticket = Ticket(
            key=f"ZL-{1001 + created}",
            type_code=type_code,
            title=f"{pop} {title} - {customer.name}",
            customer_id=customer.id,
            pop=pop,
            status="review",
            window_start=start,
            window_end=end,
            impact=impact,
            reason=ctx["reason"],
            duration="30分钟",
            subject=subject,
            body=body,
        )
        session.add(ticket)
        session.flush()
        session.add(TicketEvent(ticket_id=ticket.id, author="系统", action="创建", body="由在网业务生成的演示工单，可继续改通知正文并往下流转。"))
        created += 1
        start += timedelta(days=1)
        end += timedelta(days=1)
    return created


def main() -> None:
    parser = argparse.ArgumentParser(description="Import ZENLENET workbooks")
    parser.add_argument("--ip", required=True)
    parser.add_argument("--company", required=True)
    parser.add_argument("--ipv6", required=True)
    args = parser.parse_args()
    stats = import_all(args.ip, args.company, args.ipv6)
    for key, value in stats.items():
        print(f"{key}={value}")


if __name__ == "__main__":
    main()
