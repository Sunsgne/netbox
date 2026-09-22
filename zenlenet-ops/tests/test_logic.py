import os
import re
import tempfile
import unittest
from datetime import date
from pathlib import Path

os.environ["DATABASE_PATH"] = str(Path(tempfile.gettempdir()) / "zenlenet-ops-test.db")
os.environ["ADMIN_PASSWORD"] = "test-pass-123"
os.environ["ADMIN_USER"] = "admin"
os.environ["SECRET_KEY"] = "test-secret"
os.environ["SESSION_HTTPS_ONLY"] = "0"

from openpyxl import Workbook

from app.parsing import (
    classify_customer,
    clean_remark,
    decide_status,
    extract_ips,
    parse_bandwidth,
    render_notice,
)


class ParsingTests(unittest.TestCase):
    def test_extract_range_and_prefix(self):
        self.assertEqual(
            extract_ips("104.245.144.202-203"),
            [("104.245.144.202", 32), ("104.245.144.203", 32)],
        )
        self.assertEqual(extract_ips("216.132.130.0/24"), [("216.132.130.0", 24)])

    def test_bandwidth(self):
        self.assertEqual(parse_bandwidth("保底500M突1G"), 500)
        self.assertEqual(parse_bandwidth("2.5G"), 2500)
        self.assertIsNone(parse_bandwidth("95th"))

    def test_customer_and_secret(self):
        self.assertEqual(classify_customer("LightWAN")[1], "LightWAN")
        self.assertEqual(classify_customer("轻网预留"), ("hold", "轻网"))
        self.assertEqual(classify_customer("云森处测试")[1], "云森处")
        self.assertEqual(classify_customer("网关")[0], "infra")
        self.assertEqual(classify_customer("ser-1.szbwx-1f-h02-u5.szx1")[0], "infra")
        self.assertEqual(classify_customer("2025.10.10 ruike回收")[0], "skip")
        self.assertEqual(classify_customer("ubuntu/Secret123")[0], "skip")
        self.assertEqual(clean_remark("到期", "ubuntu/Secret123"), "到期")

    def test_returning_when_expiry_passed(self):
        status, expiry = decide_status("customer", "LightWAN", "2026.5.31退", "", date(2026, 9, 22))
        self.assertEqual(status, "returning")
        self.assertEqual(expiry, date(2026, 5, 31))

    def test_notice(self):
        text = render_notice("维护【地区/节点】，【开始时间】", {"place": "HKG", "start": "明天"})
        self.assertIn("HKG", text)
        self.assertIn("明天", text)


class AppTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        db_path = Path(os.environ["DATABASE_PATH"])
        if db_path.exists():
            db_path.unlink()
        from app.importer import import_all
        from app.main import app

        root = Path(tempfile.mkdtemp())
        ip_book = Workbook()
        sheet = ip_book.active
        sheet.title = "香港"
        sheet["A3"] = "ISP"
        sheet["B3"] = "IP（IDC）"
        sheet["C3"] = "CUSTOMER"
        sheet["A4"] = "Zenlenet"
        sheet["B4"] = "203.0.113.10"
        sheet["C4"] = "LightWAN"
        sheet["A5"] = "Zenlenet"
        sheet["B5"] = "203.0.113.11"
        sheet["C5"] = "ubuntu/Secret123"
        sheet["D5"] = "root/not-a-real-secret"
        ip_path = root / "ip.xlsx"
        ip_book.save(ip_path)

        company = Workbook()
        order = company.active
        order.title = "RM&IPT业务记录"
        order.append(["Customer", "Country", "Airport Code", "Bandwidth", "Bandwidth-add/reduce", "Type", "IP Used", "ISP"])
        order.append(["LightWAN", "香港", "HKG", "100M", None, "IPT", "203.0.113.10", "Zenlenet"])
        server = company.create_sheet("SERVER业务记录第二版")
        server.append(["Customer", "Type", "Country", "Bandwidth", "Bandwidth - Change", "Merge Bandwidth", "Specification", "Specification-Change", "IP", "IP-Change", "ISP", "Date of service provision", "Date of service end"])
        server.append(["Super", "VM", "LAX", "95th", None, None, "4c4g", None, "203.0.113.20", None, "Zenlenet", None, None])
        circuits = company.create_sheet("供应商专线记录")
        circuits.append(["CircuitID", "A_City", "A_VLAN", "A_Port", "Z_City", "Z_VLAN", "Z_Port", "Bandwidth (Mbps)", "1h Peak Traffic (Mbps)"])
        circuits.append(["dcc-test", "HKG", 100, None, "LAX", 100, None, 500, 20])
        company_path = root / "company.xlsx"
        company.save(company_path)

        ipv6 = Workbook()
        v6 = ipv6.active
        v6.title = "洛杉矶"
        v6.append(["ISP", "IP（LOCAL）", "CUSTOMER"])
        v6.append(["Zenlenet", "2602:f414:0:1::/64", "紫鸟"])
        ipv6_path = root / "ipv6.xlsx"
        ipv6.save(ipv6_path)

        stats = import_all(str(ip_path), str(company_path), str(ipv6_path))
        cls.stats = stats
        cls.app = app

    def test_import_hides_secrets(self):
        self.assertGreaterEqual(self.stats["addresses"], 2)
        self.assertGreaterEqual(self.stats["orders"], 2)
        blob = Path(os.environ["DATABASE_PATH"]).read_bytes()
        self.assertNotIn(b"Secret123", blob)
        self.assertNotIn(b"not-a-real-secret", blob)

    def test_login_billing_and_ticket(self):
        from fastapi.testclient import TestClient

        client = TestClient(self.app)
        login_page = client.get("/login")
        self.assertEqual(login_page.status_code, 200)
        token = re.search(r'name="csrf" value="([^"]+)"', login_page.text).group(1)
        blocked = client.get("/", follow_redirects=False)
        self.assertEqual(blocked.status_code, 303)
        logged = client.post(
            "/login",
            data={"username": "admin", "password": "test-pass-123", "csrf": token},
            follow_redirects=False,
        )
        self.assertEqual(logged.status_code, 303)
        home = client.get("/")
        self.assertEqual(home.status_code, 200)
        self.assertIn("LightWAN", home.text)
        self.assertIn("香港", home.text)
        self.assertNotIn("built-in method", home.text)
        summary = client.get("/resources/summary")
        self.assertIn("203.0.113.0/24", summary.text)
        self.assertIn("香港", summary.text)
        self.assertNotIn("built-in method", summary.text)
        resources = client.get("/resources?q=203.0.113.10")
        self.assertIn("203.0.113.10", resources.text)
        self.assertNotIn("Secret123", resources.text)
        billing = client.get("/billing")
        token = re.search(r'name="csrf" value="([^"]+)"', billing.text).group(1)
        generated = client.post("/billing/generate", data={"period": "2026-09", "csrf": token}, follow_redirects=True)
        self.assertIn("ZL-INV-", generated.text)
        form = client.get("/tickets/new")
        token = re.search(r'name="csrf" value="([^"]+)"', form.text).group(1)
        created = client.post(
            "/tickets",
            data={
                "csrf": token,
                "type_code": "cutover",
                "customer_id": "1",
                "pop": "HKG",
                "impact": "203.0.113.10",
                "reason": "配置升级调整割接",
                "duration": "30分钟",
                "start": "2026-10-01T01:00",
                "end": "2026-10-01T03:00",
            },
            follow_redirects=True,
        )
        self.assertIn("线路割接", created.text)
        self.assertIn("203.0.113.10", created.text)
        ticket_id = re.search(r'action="/tickets/(\d+)"', created.text).group(1)
        token = re.search(r'name="csrf" value="([^"]+)"', created.text).group(1)
        moved = client.post(
            f"/tickets/{ticket_id}",
            data={"csrf": token, "action": "next", "subject": "主题", "body": "通知正文", "comment": "评审通过"},
            follow_redirects=True,
        )
        self.assertIn("评审通过", moved.text)


if __name__ == "__main__":
    unittest.main()
